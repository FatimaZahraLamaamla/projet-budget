# -*- coding: utf-8 -*-
"""
Actualise la feuille "Donnees" du Tableau de Bord depuis le fichier source.
Regles de calcul du montant selon statut d'engagement (col AE) :
- EN COURS DE PREENGAGEMENT          : col N (MONTANT PREENGAGE INITIAL)
- PREENGAGE + statut support ATTRIBUE: col AD (TOTAL MARCHE TTC)
- PREENGAGE + autre (PHASE AOO...)   : col N (MONTANT PREENGAGE INITIAL)
- EN COURS D'ENGAGEMENT              : col AD (TOTAL MARCHE TTC)
- ENGAGE                             : col AD (TOTAL MARCHE TTC)
"""
import openpyxl, os, sys

BASE  = os.path.dirname(os.path.abspath(__file__))
SRC   = os.path.join(BASE, "ETAT D'ENGAGEMENTS JUIN 2026 FZ.xlsx")
DASH  = os.path.join(BASE, "TABLEAU_DE_BORD_ENGAGEMENTS_JUIN_2026.xlsx")

def get_montant(statut, statut_sup, montant_n, montant_ad):
    s = (statut or "").strip().upper()
    sup = (statut_sup or "").strip().upper()
    n  = montant_n  or 0
    ad = montant_ad or 0
    if s == "EN COURS DE PREENGAGEMENT":
        return n
    elif s == "PREENGAGE":
        return ad if "ATTRIBU" in sup else n
    else:
        return ad

def main():
    for f in (SRC, DASH):
        if not os.path.exists(f):
            print(f"ERREUR - Fichier introuvable : {f}")
            input("Appuyez sur Entree pour fermer...")
            sys.exit(1)

    print("Lecture du fichier source...")
    wb_src = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws_src = wb_src["ENGAGEMENTS 2026"]

    rows_data = []

    for row in ws_src.iter_rows(min_row=2, values_only=True):
        direction  = row[1]    # col B
        nature     = row[11]   # col L
        statut_sup = row[5]    # col F  - STATUT SUPPORT D'ENGAGEMENT
        montant_n  = row[13]   # col N  - MONTANT PREENGAGE INITIAL
        montant_ad = row[29]   # col AD - TOTAL MARCHE TTC
        statut     = row[30]   # col AE - STATUT D'ENGAGEMENT

        if direction is None and statut is None:
            continue

        montant = get_montant(statut, statut_sup, montant_n, montant_ad)
        rows_data.append((direction, nature, statut, montant))

    wb_src.close()
    print(f"  {len(rows_data)} lignes lues")

    print("Mise a jour du Tableau de Bord...")
    wb_dash = openpyxl.load_workbook(DASH)
    ws_don  = wb_dash["Données"]

    for row_idx in range(2, ws_don.max_row + 1):
        for col_idx in range(1, 5):
            ws_don.cell(row=row_idx, column=col_idx).value = None

    for i, (direction, nature, statut, montant) in enumerate(rows_data):
        r = i + 2
        ws_don.cell(row=r, column=1).value = direction
        ws_don.cell(row=r, column=2).value = nature
        ws_don.cell(row=r, column=3).value = statut
        ws_don.cell(row=r, column=4).value = montant

    wb_dash.save(DASH)
    total = sum(r[3] for r in rows_data if r[3])
    print(f"  {len(rows_data)} lignes ecrites")
    print(f"  Total montants : {total:,.2f} MAD")
    print(f"\nOK - Tableau de Bord mis a jour : {os.path.basename(DASH)}")

if __name__ == "__main__":
    main()
    input("\nAppuyez sur Entree pour fermer...")
