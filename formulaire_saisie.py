# -*- coding: utf-8 -*-
"""
Formulaire de saisie des engagements - SNRT 2026
Insere une nouvelle ligne dans ETAT D'ENGAGEMENTS JUIN 2026 FZ.xlsx
"""
import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl
import os
from datetime import datetime

BASE = os.path.dirname(os.path.abspath(__file__))
SRC  = os.path.join(BASE, "ETAT D'ENGAGEMENTS JUIN 2026 FZ.xlsx")

# ── Listes déroulantes ────────────────────────────────────────────────────────
DIRECTIONS = [
    "ARRIYADIA","ASSADISSA","DA","DACG","DAR","DCPA","DCRI","DG",
    "DINFORADIO","DINFOTV","DJURIDIQUE","DMPTV","DPMG","DPPR","DRH",
    "DRPSD","DTSI","MAGHRIBIA","TAMAZIGHT","TDF -SITES","TDF-EXP&MAIN","TV LAAYOUNE"
]
TYPES_ACHAT = ["AOO-M","AOO-CT","BCL","BCI","CONTRAT"]
STATUTS_ENGAGEMENT = [
    "EN COURS DE PREENGAGEMENT","PREENGAGE",
    "EN COURS D'ENGAGEMENT","ENGAGE","NON ENGAGE"
]
STATUTS_SUPPORT = ["PHASE AOO","ATTRIBUE","APPROUVE","EN SIGNATURE","INFRUCTUEUX","ANNULE"]
CATEGORIES     = ["GENERAL","PME","NON CONCERNE"]
TYPES_BUDGET   = ["Ordinaire","Exceptionnel","REPORT BUDGET ANNEE ANTERIEURE","REPORT BUDGET EXCEP  ANNEE ANTERIEURE"]
NATURES        = ["FONCTIONNEMENT","INVESTISSEMENT"]
NATURES_BESOIN = ["BESOIN RECURRENT","NOUVEAU BESOIN","RECONDUCTION","RECONDUCTIBLE"]

LIGNES_FONCT = [
    "Achats de productions et de programmes: Acquis definitive Programmes TV externes",
    "Achats de productions et de programmes: Acquis definitive pgm Radio externes",
    "Achats de productions et de programmes: Achats des programmes en co-production",
    "Achats de productions et de programmes: Achat des programmes en production interne - personnes physiques",
    "Achats de droits: Achats des droits de captation",
    "Achats de droits: Cession de Droits de sport",
    "Achats de matieres et fournitures consommables",
    "Achats non stockes de matieres et fournitures",
    "Prestations de service: Achat d'etudes",
    "Prestations de service: Achat prestations de services locales - personnes physiques",
    "Prestations de service: Achat prestations de services locales - societes",
    "Prestations de service: Achat prestations de services Etrangeres",
    "Redevances satellitaires",
    "Locations",
    "Entretien, reparation et assurance",
    "Honoraires et remuneration du personnel occasionnel",
    "Redevances et cotisations",
    "Deplacements, missions et receptions",
    "Charges marketing et relations publiques",
    "Autres charges de fonctionnement",
    "Charges du personnel",
    "Charges sociales",
    "Charges sociales diverses",
]
LIGNES_INVEST = [
    "Frais d'acquisition des immobilisations",
    "Prestations de service a immobiliser",
    "Immobilisations incorporelles (logiciels, licences, etc.)",
    "Agencements et amenagement constructions",
    "Installations techniques",
    "Equipements de telediffusion",
    "Equipements techniques",
    "Mobilier de bureau",
    "Materiel de bureau",
    "Materiel Informatique",
    "Autre Materiel et mobilier de bureau",
    "Amenagement et construction de batiments a usage technique",
    "Amenagement et construction de batiments administratifs",
    "Electrification et installations techniques",
    "Equipements audiovisuels RADIO",
    "Equipements audiovisuels TV",
    "Frais d'acquisition",
    "Habillage Radio",
    "Logiciels",
    "Materiel de diffusion",
    "Materiel de transmission",
    "Materiel informatique",
    "Moyens de contribution mobiles",
    "Moyens mobiles",
]

# ── Application ───────────────────────────────────────────────────────────────
class FormulaireEngagement:
    def __init__(self, root):
        self.root = root
        self.root.title("Saisie Engagement - SNRT 2026")
        self.root.resizable(True, True)
        self._build_ui()
        self._on_statut_change()
        self._on_nature_change()

    def _lbl(self, parent, text, row, col=0):
        tk.Label(parent, text=text, anchor="w", font=("Segoe UI", 9)).grid(
            row=row, column=col, sticky="w", padx=8, pady=3)

    def _combo(self, parent, values, row, col=1, width=40):
        cb = ttk.Combobox(parent, values=values, width=width, state="readonly", font=("Segoe UI", 9))
        cb.grid(row=row, column=col, sticky="ew", padx=8, pady=3)
        return cb

    def _entry(self, parent, row, col=1, width=42):
        e = ttk.Entry(parent, width=width, font=("Segoe UI", 9))
        e.grid(row=row, column=col, sticky="ew", padx=8, pady=3)
        return e

    def _build_ui(self):
        # ── Style ──
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("TCombobox", padding=3)
        style.configure("TEntry", padding=3)

        # ── Frame principal avec scroll ──
        canvas = tk.Canvas(self.root)
        scrollbar = ttk.Scrollbar(self.root, orient="vertical", command=canvas.yview)
        self.frame = tk.Frame(canvas, padx=10, pady=10)
        self.frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self.frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        f = self.frame
        r = 0

        # ── Titre ──
        tk.Label(f, text="NOUVEAU ENGAGEMENT", font=("Segoe UI", 12, "bold"),
                 fg="#1a5276").grid(row=r, column=0, columnspan=2, pady=(5,15))
        r += 1

        # ── Section : Identification ──
        self._section(f, "IDENTIFICATION", r); r += 1
        self._lbl(f, "Année lancement *", r)
        self.annee = self._entry(f, r, width=10); r += 1

        self._lbl(f, "Direction *", r)
        self.direction = self._combo(f, DIRECTIONS, r); r += 1

        self._lbl(f, "Référence", r)
        self.ref = self._entry(f, r); r += 1

        self._lbl(f, "Type d'achat *", r)
        self.type_achat = self._combo(f, TYPES_ACHAT, r); r += 1

        self._lbl(f, "Objet *", r)
        self.objet = self._entry(f, r, width=42); r += 1

        self._lbl(f, "Catégorie", r)
        self.categorie = self._combo(f, CATEGORIES, r); r += 1

        self._lbl(f, "Lots", r)
        self.lots = self._entry(f, r, width=10); r += 1

        # ── Section : Budget ──
        self._section(f, "BUDGET", r); r += 1

        self._lbl(f, "Nature (Budget) *", r)
        self.nature = self._combo(f, NATURES, r)
        self.nature.bind("<<ComboboxSelected>>", self._on_nature_change)
        r += 1

        self._lbl(f, "Type budget", r)
        self.type_budget = self._combo(f, TYPES_BUDGET, r); r += 1

        self._lbl(f, "Ligne budgétaire", r)
        self.ligne_budg = self._combo(f, [], r, width=40); r += 1

        self._lbl(f, "Nature besoin", r)
        self.nature_besoin = self._combo(f, NATURES_BESOIN, r); r += 1

        # ── Section : Statut & Support ──
        self._section(f, "STATUT D'ENGAGEMENT", r); r += 1

        self._lbl(f, "Statut d'engagement *", r)
        self.statut_eng = self._combo(f, STATUTS_ENGAGEMENT, r)
        self.statut_eng.bind("<<ComboboxSelected>>", self._on_statut_change)
        r += 1

        self._lbl(f, "Statut 1", r)
        self.statut1 = self._entry(f, r, width=20); r += 1

        self._lbl(f, "Statut support engagement", r)
        self.statut_sup = self._combo(f, STATUTS_SUPPORT, r)
        self.statut_sup.bind("<<ComboboxSelected>>", self._on_statut_change)
        r += 1

        self._lbl(f, "N° Support d'engagement", r)
        self.num_support = self._entry(f, r); r += 1

        # ── Section : Montants ──
        self._section(f, "MONTANTS", r); r += 1

        # Explication dynamique
        self.lbl_montant_info = tk.Label(f, text="", fg="#1a5276",
                                          font=("Segoe UI", 9, "italic"), wraplength=400, justify="left")
        self.lbl_montant_info.grid(row=r, column=0, columnspan=2, sticky="w", padx=8, pady=(0,5))
        r += 1

        # Montant préengagé initial
        self.lbl_preengage = tk.Label(f, text="Montant préengagé initial (MAD)", anchor="w", font=("Segoe UI", 9))
        self.lbl_preengage.grid(row=r, column=0, sticky="w", padx=8, pady=3)
        self.montant_n = ttk.Entry(f, width=20, font=("Segoe UI", 9))
        self.montant_n.grid(row=r, column=1, sticky="w", padx=8, pady=3)
        self.row_preengage = r; r += 1

        # Estimation TTC
        self._lbl(f, "Estimation TTC (MAD)", r)
        self.estimation_ttc = self._entry(f, r, width=20); r += 1

        # Estimation HT
        self._lbl(f, "Estimation HT (MAD)", r)
        self.estimation_ht = self._entry(f, r, width=20); r += 1

        # Total Marché TTC
        self.lbl_total_ttc = tk.Label(f, text="Total Marché TTC / Engagé TTC (MAD)", anchor="w", font=("Segoe UI", 9))
        self.lbl_total_ttc.grid(row=r, column=0, sticky="w", padx=8, pady=3)
        self.montant_ad = ttk.Entry(f, width=20, font=("Segoe UI", 9))
        self.montant_ad.grid(row=r, column=1, sticky="w", padx=8, pady=3)
        self.row_total_ttc = r; r += 1

        # ── Section : Adjudicataire & Dates ──
        self._section(f, "ADJUDICATAIRE & DATES", r); r += 1

        self._lbl(f, "Adjudicataire / Fournisseur", r)
        self.adjudicataire = self._entry(f, r); r += 1

        self._lbl(f, "Année approbation", r)
        self.annee_appro = self._entry(f, r, width=10); r += 1

        self._lbl(f, "Date de la fiche (JJ/MM/AAAA)", r)
        self.date_fiche = self._entry(f, r, width=15); r += 1

        self._lbl(f, "Date visa (JJ/MM/AAAA)", r)
        self.date_visa = self._entry(f, r, width=15); r += 1

        self._lbl(f, "Date approbation (JJ/MM/AAAA)", r)
        self.date_approb = self._entry(f, r, width=15); r += 1

        self._lbl(f, "Remarque", r)
        self.remarque = self._entry(f, r); r += 1

        # ── Boutons ──
        btn_frame = tk.Frame(f)
        btn_frame.grid(row=r, column=0, columnspan=2, pady=15)
        ttk.Button(btn_frame, text="  Ajouter la ligne  ", command=self._ajouter).pack(side="left", padx=10)
        ttk.Button(btn_frame, text="  Effacer  ", command=self._effacer).pack(side="left", padx=10)

        # ── Status bar ──
        self.status = tk.Label(self.root, text="Prêt.", anchor="w", fg="gray", font=("Segoe UI", 9))
        self.status.pack(fill="x", padx=10, pady=4)

    def _section(self, parent, title, row):
        tk.Label(parent, text=title, font=("Segoe UI", 9, "bold"),
                 fg="white", bg="#1a5276", anchor="w", padx=6).grid(
            row=row, column=0, columnspan=2, sticky="ew", pady=(10, 2))

    def _on_nature_change(self, event=None):
        nature = self.nature.get()
        if nature == "FONCTIONNEMENT":
            self.ligne_budg["values"] = LIGNES_FONCT
        elif nature == "INVESTISSEMENT":
            self.ligne_budg["values"] = LIGNES_INVEST
        else:
            self.ligne_budg["values"] = LIGNES_FONCT + LIGNES_INVEST
        self.ligne_budg.set("")

    def _on_statut_change(self, event=None):
        statut = self.statut_eng.get().upper()
        sup    = self.statut_sup.get().upper()

        # Couleurs indicatives
        COLOR_ACTIVE   = "#1a5276"
        COLOR_INACTIVE = "#aaaaaa"

        if statut == "EN COURS DE PREENGAGEMENT":
            self._show_montant(preengage=True, total_ttc=False)
            info = "→ Saisir le MONTANT PRÉENGAGÉ INITIAL (colonne N)"
        elif statut == "PREENGAGE":
            if "ATTRIBU" in sup:
                self._show_montant(preengage=False, total_ttc=True)
                info = "→ Statut support ATTRIBUÉ : saisir le TOTAL MARCHÉ TTC (colonne AD)"
            else:
                self._show_montant(preengage=True, total_ttc=False)
                info = "→ Statut support PHASE AOO : saisir le MONTANT PRÉENGAGÉ INITIAL (colonne N)"
        elif statut in ("EN COURS D'ENGAGEMENT", "ENGAGE"):
            self._show_montant(preengage=False, total_ttc=True)
            info = "→ Saisir le TOTAL MARCHÉ TTC / ENGAGÉ TTC (colonne AD)"
        else:
            self._show_montant(preengage=False, total_ttc=False)
            info = "→ Aucun montant requis pour ce statut"

        self.lbl_montant_info.config(text=info)

    def _show_montant(self, preengage, total_ttc):
        fg_n  = "#1a5276" if preengage  else "#aaaaaa"
        fg_ad = "#1a5276" if total_ttc  else "#aaaaaa"
        state_n  = "normal" if preengage  else "disabled"
        state_ad = "normal" if total_ttc  else "disabled"
        self.lbl_preengage.config(fg=fg_n)
        self.lbl_total_ttc.config(fg=fg_ad)
        self.montant_n.config(state=state_n)
        self.montant_ad.config(state=state_ad)
        if not preengage:
            self.montant_n.delete(0, "end")
        if not total_ttc:
            self.montant_ad.delete(0, "end")

    def _parse_date(self, val):
        val = val.strip()
        if not val:
            return None
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(val, fmt)
            except ValueError:
                continue
        return None

    def _parse_float(self, val):
        val = val.strip().replace(" ", "").replace(",", ".")
        if not val:
            return None
        try:
            return float(val)
        except ValueError:
            return None

    def _effacer(self):
        for widget in [self.annee, self.ref, self.objet, self.lots,
                       self.statut1, self.num_support, self.montant_n,
                       self.estimation_ttc, self.estimation_ht, self.montant_ad,
                       self.adjudicataire, self.annee_appro, self.date_fiche,
                       self.date_visa, self.date_approb, self.remarque]:
            widget.config(state="normal")
            widget.delete(0, "end")
        for cb in [self.direction, self.type_achat, self.categorie, self.nature,
                   self.type_budget, self.ligne_budg, self.nature_besoin,
                   self.statut_eng, self.statut_sup]:
            cb.set("")
        self._on_statut_change()
        self.status.config(text="Formulaire effacé.", fg="gray")

    def _ajouter(self):
        # ── Validation champs obligatoires ──
        erreurs = []
        if not self.annee.get().strip():       erreurs.append("Année lancement")
        if not self.direction.get():           erreurs.append("Direction")
        if not self.type_achat.get():          erreurs.append("Type d'achat")
        if not self.objet.get().strip():       erreurs.append("Objet")
        if not self.statut_eng.get():          erreurs.append("Statut d'engagement")
        if not self.nature.get():              erreurs.append("Nature (Budget)")
        if erreurs:
            messagebox.showwarning("Champs obligatoires manquants",
                                   "Veuillez renseigner :\n- " + "\n- ".join(erreurs))
            return

        # ── Collecte des valeurs ──
        statut  = self.statut_eng.get().upper()
        sup     = self.statut_sup.get().upper()
        mn      = self._parse_float(self.montant_n.get())
        mad     = self._parse_float(self.montant_ad.get())

        # Vérification cohérence montant
        if statut == "EN COURS DE PREENGAGEMENT" and mn is None:
            if not messagebox.askyesno("Montant manquant",
                "Le montant préengagé initial est vide. Continuer quand même ?"):
                return
        if statut in ("EN COURS D'ENGAGEMENT", "ENGAGE") and mad is None:
            if not messagebox.askyesno("Montant manquant",
                "Le Total Marché TTC est vide. Continuer quand même ?"):
                return

        # ── Écriture dans le fichier Excel ──
        try:
            wb = openpyxl.load_workbook(SRC)
            wb.calculation.fullCalcOnLoad = True
            ws = wb["ENGAGEMENTS 2026"]

            # Trouver la première ligne vide
            next_row = ws.max_row + 1
            for r in range(ws.max_row, 1, -1):
                if any(ws.cell(row=r, column=c).value for c in range(1, 40)):
                    next_row = r + 1
                    break

            def w(col, val):
                if val is not None and val != "":
                    ws.cell(row=next_row, column=col).value = val

            w(1,  self._parse_float(self.annee.get()) or self.annee.get().strip())  # A
            w(2,  self.direction.get())          # B
            w(3,  self.ref.get().strip())        # C
            w(4,  self.type_achat.get())         # D
            w(5,  self.statut1.get().strip())    # E
            w(6,  self.statut_sup.get())         # F
            w(7,  self.num_support.get().strip())# G
            w(8,  self.categorie.get())          # H
            w(9,  self.lots.get().strip())       # I
            w(10, self.objet.get().strip())      # J
            w(11, self.type_budget.get())        # K
            w(12, self.nature.get())             # L
            w(13, self.ligne_budg.get())         # M
            w(14, mn)                            # N  MONTANT PREENGAGE INITIAL
            w(15, self._parse_float(self.estimation_ttc.get()))  # O
            w(16, self._parse_float(self.estimation_ht.get()))   # P
            w(17, self.nature_besoin.get())      # Q
            w(20, self.adjudicataire.get().strip()) # T
            w(21, self._parse_float(self.annee_appro.get()) or self.annee_appro.get().strip())  # U
            w(22, self._parse_date(self.date_fiche.get()))   # V
            w(23, self._parse_date(self.date_visa.get()))    # W
            w(24, self._parse_date(self.date_approb.get()))  # X
            w(30, mad)                           # AD TOTAL MARCHE TTC
            w(31, self.statut_eng.get())         # AE STATUT D'ENGAGEMENT
            w(39, self.remarque.get().strip())   # AM

            wb.save(SRC)
            wb.close()

            msg = f"Ligne {next_row} ajoutée avec succès dans le fichier source."
            self.status.config(text=msg, fg="#1a5276")
            messagebox.showinfo("Succès", msg + "\n\nN'oubliez pas d'actualiser le tableau de bord.")
            self._effacer()

        except PermissionError:
            messagebox.showerror("Fichier ouvert",
                "Le fichier Excel est ouvert. Fermez-le d'abord puis réessayez.")
        except Exception as ex:
            messagebox.showerror("Erreur", str(ex))


if __name__ == "__main__":
    root = tk.Tk()
    app = FormulaireEngagement(root)
    root.mainloop()
