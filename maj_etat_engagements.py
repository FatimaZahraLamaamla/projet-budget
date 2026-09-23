"""
Mise à jour du classeur de tableaux de bord ETAT_D_ENGAGEMENTS.xlsx
à partir d'un nouvel état d'engagements source.

Les onglets sources (ENGAGEMENTS 2026, Parametres, lignes) sont remplacés
par ceux du fichier source ; les onglets de tableaux de bord (Tableau de Bord,
Filtre Interactif, Tableau de Bord Suivi, Données Suivi, PPM 2026) sont
conservés tels quels et se recalculent à l'ouverture dans Excel.
L'onglet ENGAGEMENTS 2026 reçoit en colonne A une icône d'état (les colonnes
source sont décalées d'un cran, les formules des tableaux de bord aussi) et
est remis en forme (voir style_engagements.py).

Usage :
    python maj_etat_engagements.py <fichier_source.xlsx> [modele.xlsx] [sortie.xlsx]

Par défaut : modele = ETAT_D_ENGAGEMENTS.xlsx,
             sortie = ETAT_D_ENGAGEMENTS - MAJ <date>.xlsx
"""
import sys
import warnings
from copy import copy
from datetime import date

import openpyxl
from openpyxl.worksheet.table import Table

from decalage_colonnes import decaler_formule, decaler_reference
from style_engagements import ENTETE_ETAT, appliquer_style

warnings.filterwarnings("ignore", category=UserWarning)

ONGLETS_SOURCES = ["Parametres", "ENGAGEMENTS 2026", "lignes"]
ONGLET_ENGAGEMENTS = "ENGAGEMENTS 2026"
LIGNES_FORMULES_SUIVI = 3000  # plage couverte par l'onglet "Données Suivi"


def derniere_ligne(ws, col):
    """Dernière ligne dont la colonne DIRECTION est renseignée."""
    fin = 1
    for r, (v,) in enumerate(
            ws.iter_rows(min_row=2, min_col=col, max_col=col, values_only=True), start=2):
        if v not in (None, ""):
            fin = r
    return fin


def copier_onglet(src, dst, decalage=0, entetes=()):
    """Copie un onglet d'un classeur à l'autre. Avec decalage=n, les colonnes
    sont décalées de n vers la droite (formules, validations, largeurs,
    tableau et volets figés suivent) pour libérer les n premières colonnes,
    dont les en-têtes sont donnés par `entetes`."""
    def dec(formule):
        return decaler_formule(formule, dst.title, dst.title, decalage) if decalage else formule

    for row in src.iter_rows():
        for c in row:
            d = dst.cell(row=c.row, column=c.column + decalage, value=dec(c.value))
            if c.has_style:
                d.font = copy(c.font)
                d.fill = copy(c.fill)
                d.border = copy(c.border)
                d.alignment = copy(c.alignment)
                d.number_format = c.number_format
                d.protection = copy(c.protection)
            if c.hyperlink:
                d.hyperlink = copy(c.hyperlink)
            if c.comment:
                d.comment = copy(c.comment)

    for key, dim in src.column_dimensions.items():
        dd = dst.column_dimensions[decaler_reference(f"{key}1", decalage)[:-1]]
        dd.width, dd.hidden = dim.width, dim.hidden
        dd.min, dd.max = dim.min + decalage, dim.max + decalage
        dd.outlineLevel = dim.outlineLevel
    for key, dim in src.row_dimensions.items():
        dd = dst.row_dimensions[key]
        dd.height, dd.hidden = dim.height, dim.hidden
        dd.outlineLevel = dim.outlineLevel

    for rng in src.merged_cells.ranges:
        dst.merge_cells(decaler_reference(str(rng), decalage))
    if not decalage:  # avec décalage, la mise en forme conditionnelle est refaite par le style
        for cf in src.conditional_formatting:
            for rule in cf.rules:
                dst.conditional_formatting.add(str(cf.sqref), rule)
    for dv in src.data_validations.dataValidation:
        ndv = copy(dv)
        if decalage:
            ndv.sqref = decaler_reference(dv.sqref, decalage)
            for attr in ("formula1", "formula2"):
                f = getattr(dv, attr)
                if f:
                    setattr(ndv, attr, dec("=" + f)[1:])
        dst.add_data_validation(ndv)
    for img in getattr(src, "_images", []):
        dst.add_image(img)

    # Volets figés : on reprend le découpage réel (xSplit/ySplit) et non la
    # position de défilement (topLeftCell), qu'openpyxl confond avec le figeage.
    pane = src.sheet_view.pane
    if pane is not None and pane.state == "frozen":
        dst.freeze_panes = dst.cell(row=int(pane.ySplit or 0) + 1,
                                    column=int(pane.xSplit or 0) + 1 + decalage)
    if src.auto_filter.ref:
        dst.auto_filter.ref = decaler_reference(src.auto_filter.ref, decalage)
    dst.sheet_properties.tabColor = src.sheet_properties.tabColor
    dst.sheet_view.zoomScale = src.sheet_view.zoomScale
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.page_setup.orientation = src.page_setup.orientation
    dst.page_setup.paperSize = src.page_setup.paperSize

    for t in src.tables.values():
        noms = [c.name for c in t.tableColumns]
        ref = t.ref
        if decalage:
            # le tableau englobe les colonnes libérées en tête
            debut, fin = decaler_reference(t.ref, decalage).split(":")
            ref = f"A{debut.lstrip('$ABCDEFGHIJKLMNOPQRSTUVWXYZ')}:{fin}"
            noms = list(entetes) + noms
            for i, nom in enumerate(entetes, start=1):
                dst.cell(row=int(debut.lstrip("$ABCDEFGHIJKLMNOPQRSTUVWXYZ")), column=i, value=nom)
        nt = Table(displayName=t.displayName, ref=ref)
        nt.tableStyleInfo = copy(t.tableStyleInfo)
        nt.autoFilter = copy(t.autoFilter)
        if nt.autoFilter is not None:
            nt.autoFilter.ref = ref
            nt.autoFilter.filterColumn = []
        nt._initialise_columns()
        for nom, col_dst in zip(noms, nt.tableColumns):
            col_dst.name = nom
        dst.add_table(nt)


def decaler_references_classeur(wb, onglet, n):
    """Décale de n colonnes toutes les références vers `onglet` situées dans
    les autres onglets (formules) et dans les noms définis."""
    nb = 0
    for ws in wb.worksheets:
        if ws.title == onglet:
            continue
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and onglet in v:
                    nv = decaler_formule(v, onglet, ws.title, n)
                    if nv != v:
                        c.value = nv
                        nb += 1
    for nom, dn in wb.defined_names.items():
        if dn.attr_text and onglet in dn.attr_text:
            dn.attr_text = decaler_formule("=" + dn.attr_text, onglet, "", n)[1:]
            nb += 1
    return nb


def retirer_filtres(ws):
    """Supprime les filtres enregistrés dans le fichier source et réaffiche
    les lignes qu'ils masquaient : openpyxl ne réapplique pas les filtres,
    les lignes ajoutées resteraient visibles et les autres masquées."""
    nb = 0
    for t in ws.tables.values():
        if t.autoFilter is not None:
            nb += len(t.autoFilter.filterColumn)
            t.autoFilter.filterColumn = []
            t.autoFilter.sortState = None
    if ws.auto_filter.ref:
        nb += len(ws.auto_filter.filterColumn)
        ws.auto_filter.filterColumn = []
    masquees = 0
    for dim in ws.row_dimensions.values():
        if dim.hidden:
            dim.hidden = False
            masquees += 1
    return nb, masquees


def etendre_table(ws, fin):
    """Étend le tableau Excel de l'onglet jusqu'à la dernière ligne de données."""
    for t in ws.tables.values():
        debut, _ = t.ref.split(":")
        col_fin = "".join(ch for ch in t.ref.split(":")[1] if ch.isalpha())
        ancien = t.ref
        t.ref = f"{debut}:{col_fin}{fin}"
        if t.autoFilter is not None:
            t.autoFilter.ref = t.ref
        return ancien, t.ref
    return None, None


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    source = sys.argv[1]
    modele = sys.argv[2] if len(sys.argv) > 2 else "ETAT_D_ENGAGEMENTS.xlsx"
    sortie = sys.argv[3] if len(sys.argv) > 3 else (
        f"ETAT_D_ENGAGEMENTS - MAJ {date.today():%d-%m-%Y}.xlsx")

    wb_src = openpyxl.load_workbook(source)
    wb = openpyxl.load_workbook(modele)

    manquants = [n for n in ONGLETS_SOURCES if n not in wb_src.sheetnames]
    if manquants:
        sys.exit(f"Onglets absents du fichier source : {manquants}")

    # Colonne ÉTAT (icônes) : déjà présente dans le modèle et/ou la source ?
    modele_a_etat = wb[ONGLET_ENGAGEMENTS]["A1"].value == ENTETE_ETAT
    source_a_etat = wb_src[ONGLET_ENGAGEMENTS]["A1"].value == ENTETE_ETAT

    # Contrôle des en-têtes : la structure doit être identique au modèle
    h_mod = [c.value for c in wb[ONGLET_ENGAGEMENTS][1]][int(modele_a_etat):]
    h_src = [c.value for c in wb_src[ONGLET_ENGAGEMENTS][1]][int(source_a_etat):]
    if h_mod[:39] != h_src[:39]:
        diff = [(i + 1, a, b) for i, (a, b) in enumerate(zip(h_mod, h_src)) if a != b]
        sys.exit(f"Colonnes différentes entre modèle et source : {diff}")

    for nom in ONGLETS_SOURCES:
        idx = wb.sheetnames.index(nom)
        wb.remove(wb[nom])
        dst = wb.create_sheet(nom, idx)
        decalage = 1 if nom == ONGLET_ENGAGEMENTS and not source_a_etat else 0
        copier_onglet(wb_src[nom], dst, decalage, entetes=[ENTETE_ETAT][:decalage])

    # Les tableaux de bord du modèle lisent l'ancienne disposition : on décale
    # leurs références d'une colonne (une seule fois, au passage à la colonne ÉTAT)
    nb_refs = 0
    if not modele_a_etat:
        nb_refs = decaler_references_classeur(wb, ONGLET_ENGAGEMENTS, 1)

    ws = wb[ONGLET_ENGAGEMENTS]
    fin = derniere_ligne(ws, col=3)
    if fin > LIGNES_FORMULES_SUIVI:
        sys.exit(f"{fin} lignes : dépasse la plage de 'Données Suivi' "
                 f"({LIGNES_FORMULES_SUIVI}), à étendre.")
    ancien, nouveau = etendre_table(ws, fin)
    nb_filtres, nb_masquees = retirer_filtres(ws)
    appliquer_style(ws, fin)

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # recalcul complet à l'ouverture
    wb.save(sortie)

    print(f"Source      : {source}")
    print(f"Modèle      : {modele}")
    print(f"Sortie      : {sortie}")
    print(f"Lignes      : {fin - 1} dossiers (jusqu'à la ligne {fin})")
    if ancien:
        print(f"Tableau     : {ancien} -> {nouveau}")
    print(f"Volets figés: {ws.freeze_panes}")
    if nb_refs:
        print(f"Colonne ÉTAT: insérée en A, {nb_refs} formule(s)/nom(s) du modèle décalé(s)")
    if nb_filtres or nb_masquees:
        print(f"Filtres     : {nb_filtres} filtre(s) retiré(s), "
              f"{nb_masquees} ligne(s) réaffichée(s)")


if __name__ == "__main__":
    main()
